"""Extraktion von reinem Text aus .txt-, .docx- und .pdf-Dateien.

Die Extraktion beschränkt sich auf den sichtbaren Fließtext. Formatierungen,
Kopf- und Fußzeilen von PDFs sowie Textfelder außerhalb des normalen
Textflusses werden je nach Quellformat unterschiedlich gut erfasst; siehe
METHODIK.md.
"""

from __future__ import annotations

from pathlib import Path


class ExtractionError(RuntimeError):
    """Wird ausgelöst, wenn eine Datei nicht gelesen werden konnte."""


def _extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - hängt von der Installation ab
        raise ExtractionError(
            "Für .docx-Dateien wird das Paket python-docx benötigt "
            "(pip install python-docx)."
        ) from exc

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - hängt von der Installation ab
        raise ExtractionError(
            "Für .pdf-Dateien wird das Paket pypdf benötigt (pip install pypdf)."
        ) from exc

    reader = PdfReader(str(path))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(parts)


EXTRACTORS = {
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
}


def extract_text(path: str | Path) -> str:
    """Liest path und liefert den enthaltenen Text als eine Zeichenkette.

    Unterstützt werden .txt, .md, .docx und .pdf anhand der Dateiendung.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise ExtractionError(f"Datei nicht gefunden: {file_path}")
    suffix = file_path.suffix.lower()
    extractor = EXTRACTORS.get(suffix)
    if extractor is None:
        supported = ", ".join(sorted(EXTRACTORS))
        raise ExtractionError(
            f"Nicht unterstütztes Dateiformat '{suffix}'. Unterstützt: {supported}."
        )
    text = extractor(file_path)
    if not text.strip():
        raise ExtractionError(
            f"Aus {file_path} konnte kein Text extrahiert werden (leere oder "
            f"gescannte Datei ohne Texterkennung?)."
        )
    return text
